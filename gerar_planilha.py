#!/usr/bin/env python3
"""
Planilha Financeira Familiar — Gerador
Dados do casal: Ele R$3.600 + Ela R$2.500 + Extra R$500 = R$6.600/mês
"""
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.chart import BarChart, PieChart, Reference
import datetime

wb = openpyxl.Workbook()
wb.remove(wb.active)

# =============================================
# PALETA DE CORES
# =============================================
AZUL_ESC  = "1A5276"
AZUL_MED  = "2980B9"
AZUL_CLA  = "D6EAF8"
VERDE_ESC = "1E8449"
VERDE_MED = "27AE60"
VERDE_CLA = "D5F5E3"
VERM_ESC  = "C0392B"
VERM_CLA  = "FADBD8"
LARAN     = "D35400"
LARAN_MED = "E67E22"
LARAN_CLA = "FDEBD0"
AMAR_CLA  = "FEF9E7"
CINZA_CLA = "F2F3F4"
CINZA_MED = "BDC3C7"
ROXO_ESC  = "6C3483"
ROXO_CLA  = "E8DAEF"
BRANCO    = "FFFFFF"
PRETO     = "1C2833"

def f(color):
    return PatternFill(fill_type='solid', fgColor=color)

def font(bold=False, color=PRETO, size=10, italic=False):
    return Font(bold=bold, color=color, size=size, italic=italic, name='Calibri')

def aln(h='center', v='center', wrap=True):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

def bdr(color=CINZA_MED, style='thin'):
    s = Side(style=style, color=color)
    return Border(left=s, right=s, top=s, bottom=s)

def bdr_thick(color=AZUL_ESC):
    s = Side(style='medium', color=color)
    return Border(left=s, right=s, top=s, bottom=s)

def W(ws, cols):
    for i, w in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

def H(ws, rows: dict):
    for r, h in rows.items():
        ws.row_dimensions[r].height = h

def c(ws, row, col, value=None, bg=None, fg=PRETO, size=10, bold=False,
      ha='center', va='center', wrap=True, fmt=None, italic=False, border=None):
    cell = ws.cell(row=row, column=col)
    if value is not None:
        cell.value = value
    if bg:
        cell.fill = f(bg)
    cell.font = Font(bold=bold, color=fg, size=size, italic=italic, name='Calibri')
    cell.alignment = Alignment(horizontal=ha, vertical=va, wrap_text=wrap)
    if fmt:
        cell.number_format = fmt
    if border:
        cell.border = border
    return cell

def m(ws, r1, c1, r2, c2):
    ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)

def borders(ws, r1, c1, r2, c2, color=CINZA_MED, style='thin'):
    b = bdr(color, style)
    for row in ws.iter_rows(min_row=r1, min_col=c1, max_row=r2, max_col=c2):
        for cell in row:
            cell.border = b

def card(ws, label_row, val_row, c1, c2, label, val_formula,
         label_bg, val_bg, val_fg, val_size=18, val_fmt='R$ #,##0.00'):
    m(ws, label_row, c1, label_row, c2)
    cx = ws.cell(row=label_row, column=c1, value=label)
    cx.fill = f(label_bg)
    cx.font = Font(bold=True, color=BRANCO, size=9, name='Calibri')
    cx.alignment = Alignment(horizontal='center', vertical='center')

    m(ws, val_row, c1, val_row, c2)
    cv = ws.cell(row=val_row, column=c1, value=val_formula)
    cv.fill = f(val_bg)
    cv.font = Font(bold=True, color=val_fg, size=val_size, name='Calibri')
    cv.alignment = Alignment(horizontal='center', vertical='center')
    cv.number_format = val_fmt
    cv.border = bdr_thick(label_bg)
    return cv

today = datetime.date.today()

# =============================================
# ABA 1 — DASHBOARD
# =============================================
ws_d = wb.create_sheet("Dashboard", 0)
ws_d.sheet_view.showGridLines = False
ws_d.sheet_properties.tabColor = AZUL_ESC

W(ws_d, [1.5, 16, 16, 16, 16, 16, 16, 1.5])
H(ws_d, {1:8, 2:55, 3:8, 4:22, 5:60, 6:22, 7:55, 8:22, 9:55,
         10:12, 11:28, 12:12, 13:28, 14:12, 15:22, 16:8})

# Título principal
m(ws_d, 2, 2, 2, 7)
cx = ws_d['B2']
cx.value = "CONTROLE FINANCEIRO FAMILIAR"
cx.fill = f(AZUL_ESC)
cx.font = Font(bold=True, color=BRANCO, size=22, name='Calibri')
cx.alignment = Alignment(horizontal='center', vertical='center')

# Subtítulo data
m(ws_d, 3, 2, 3, 7)
cx = ws_d['B3']
cx.value = f"Atualizado em: {today.strftime('%d/%m/%Y')}"
cx.fill = f(AZUL_MED)
cx.font = Font(color=BRANCO, size=9, italic=True, name='Calibri')
cx.alignment = Alignment(horizontal='center', vertical='center')

# Labels row (row 4)
card_configs = [
    (4, 5, 2, 3, "ENTRADAS DO MES",
     '=SUMPRODUCT((MONTH(Lancamentos!$A$3:$A$3000)=MONTH(TODAY()))*(YEAR(Lancamentos!$A$3:$A$3000)=YEAR(TODAY()))*(Lancamentos!$H$3:$H$3000="Receita")*(Lancamentos!$G$3:$G$3000))',
     VERDE_ESC, VERDE_CLA, VERDE_ESC, 18, 'R$ #,##0.00'),
    (4, 5, 4, 5, "SAIDAS DO MES",
     '=SUMPRODUCT((MONTH(Lancamentos!$A$3:$A$3000)=MONTH(TODAY()))*(YEAR(Lancamentos!$A$3:$A$3000)=YEAR(TODAY()))*(Lancamentos!$H$3:$H$3000="Despesa")*(Lancamentos!$G$3:$G$3000))',
     VERM_ESC, VERM_CLA, VERM_ESC, 18, 'R$ #,##0.00'),
    (4, 5, 6, 7, "ECONOMIA DO MES",
     '=B5-D5',
     AZUL_MED, AZUL_CLA, AZUL_ESC, 18, 'R$ #,##0.00'),
]

for (lr, vr, c1, c2, lbl, formula, lbg, vbg, vfg, vsz, vfmt) in card_configs:
    card(ws_d, lr, vr, c1, c2, lbl, formula, lbg, vbg, vfg, vsz, vfmt)

# Row 2 de KPIs (rows 6-7)
card(ws_d, 6, 7, 2, 3, "TAXA DE POUPANCA",
     '=IF(B5=0,0,F5/B5)',
     ROXO_ESC, ROXO_CLA, ROXO_ESC, 22, '0.0%')

card(ws_d, 6, 7, 4, 5, "RESERVA ACUMULADA",
     "=Reserva!B7",
     ROXO_ESC, ROXO_CLA, ROXO_ESC, 18, 'R$ #,##0.00')

# Saúde financeira
m(ws_d, 6, 6, 6, 7)
cx = ws_d['F6']
cx.value = "SAUDE FINANCEIRA"
cx.fill = f(ROXO_ESC)
cx.font = Font(bold=True, color=BRANCO, size=9, name='Calibri')
cx.alignment = Alignment(horizontal='center', vertical='center')

m(ws_d, 7, 6, 7, 7)
cx = ws_d['F7']
cx.value = ('=IF(B5=0,"---",IF(F5/B5>=0.2,"OTIMO - Acima de 20%",'
            'IF(F5/B5>=0.10,"BOM - Entre 10% e 20%",'
            'IF(F5/B5>=0,"ATENCAO - Abaixo de 10%","CRITICO - Gastando mais do que ganha"))))')
cx.fill = f(ROXO_CLA)
cx.font = Font(bold=True, color=ROXO_ESC, size=12, name='Calibri')
cx.alignment = Alignment(horizontal='center', vertical='center')
cx.border = bdr_thick(ROXO_ESC)

# Separador
m(ws_d, 8, 2, 8, 7)
ws_d['B8'].fill = f(AZUL_ESC)

# Progresso da reserva
m(ws_d, 9, 2, 9, 7)
cx = ws_d['B9']
cx.value = "PROGRESSO DA RESERVA DE EMERGENCIA"
cx.fill = f(AZUL_MED)
cx.font = Font(bold=True, color=BRANCO, size=11, name='Calibri')
cx.alignment = Alignment(horizontal='center', vertical='center')

m(ws_d, 10, 2, 10, 7)
ws_d['B10'].fill = f(CINZA_CLA)
ws_d['B10'].value = "Valor acumulado:"
ws_d['B10'].font = Font(color=AZUL_ESC, size=9, name='Calibri')
ws_d['B10'].alignment = Alignment(horizontal='left', vertical='center')

m(ws_d, 11, 2, 11, 7)
cx = ws_d['B11']
cx.value = ('=IFERROR("R$ "&TEXT(Reserva!B7,"#.##0,00")&"  de  R$ "&TEXT(Reserva!B6,"#.##0,00")'
            '&"     |     "&TEXT(Reserva!B9,"0.0%")&" concluido  |  "& TEXT(Reserva!B10,"0.0")&" meses protegidos","---")')
cx.fill = f(VERDE_CLA)
cx.font = Font(bold=True, color=VERDE_ESC, size=13, name='Calibri')
cx.alignment = Alignment(horizontal='center', vertical='center')
cx.border = bdr_thick(VERDE_ESC)

# Dívidas e metas
m(ws_d, 12, 2, 12, 7)
ws_d['B12'].fill = f(AZUL_ESC)

m(ws_d, 13, 2, 13, 4)
ws_d['B13'].value = "FINANCIAMENTO CARRO"
ws_d['B13'].fill = f(LARAN_CLA)
ws_d['B13'].font = Font(bold=True, color=LARAN, size=10, name='Calibri')
ws_d['B13'].alignment = Alignment(horizontal='center', vertical='center')
ws_d['B13'].border = bdr_thick(LARAN)

m(ws_d, 13, 5, 13, 7)
cx = ws_d['E13']
cx.value = "=Dividas!C4&\"  |  R$ \"&TEXT(Dividas!E4,\"#.##0,00\")&\"/mes\""
cx.fill = f(LARAN_CLA)
cx.font = Font(bold=True, color=LARAN, size=10, name='Calibri')
cx.alignment = Alignment(horizontal='center', vertical='center')
cx.border = bdr_thick(LARAN)

# Meta imóvel
m(ws_d, 14, 2, 14, 7)
ws_d['B14'].fill = f(CINZA_CLA)

m(ws_d, 15, 2, 15, 4)
ws_d['B15'].value = "META: IMOVEL PROPRIO"
ws_d['B15'].fill = f(AZUL_CLA)
ws_d['B15'].font = Font(bold=True, color=AZUL_ESC, size=10, name='Calibri')
ws_d['B15'].alignment = Alignment(horizontal='center', vertical='center')
ws_d['B15'].border = bdr_thick(AZUL_ESC)

m(ws_d, 15, 5, 15, 7)
cx = ws_d['E15']
cx.value = ("=IFERROR(\"R$ \"&TEXT(Metas!E5,\"#.##0,00\")&\" acumulado  |  Prazo: \"&Metas!D5,\"Defina o valor\")")
cx.fill = f(AZUL_CLA)
cx.font = Font(bold=True, color=AZUL_ESC, size=10, name='Calibri')
cx.alignment = Alignment(horizontal='center', vertical='center')
cx.border = bdr_thick(AZUL_ESC)

m(ws_d, 16, 2, 16, 7)
cx = ws_d['B16']
cx.value = "  Registre gastos na aba LANCAMENTOS toda semana. Atualize a Reserva quando fizer um deposito."
cx.fill = f(AMAR_CLA)
cx.font = Font(color=LARAN, size=9, italic=True, name='Calibri')
cx.alignment = Alignment(horizontal='left', vertical='center')

# =============================================
# ABA 2 — LANÇAMENTOS
# =============================================
ws_l = wb.create_sheet("Lancamentos", 1)
ws_l.sheet_view.showGridLines = False
ws_l.sheet_properties.tabColor = AZUL_MED

W(ws_l, [12, 30, 18, 22, 13, 18, 15, 12])
H(ws_l, {1:14, 2:36})

# Título
m(ws_l, 1, 1, 1, 8)
cx = ws_l['A1']
cx.value = "LANCAMENTOS — Registre aqui toda entrada e saida de dinheiro"
cx.fill = f(AZUL_MED)
cx.font = Font(bold=True, color=BRANCO, size=13, name='Calibri')
cx.alignment = Alignment(horizontal='center', vertical='center')

# Cabeçalhos das colunas
headers = ["Data", "Descricao", "Categoria", "Subcategoria",
           "Responsavel", "Pagamento", "Valor (R$)", "Tipo"]
for i, h in enumerate(headers, 1):
    cx = ws_l.cell(row=2, column=i, value=h)
    cx.fill = f(AZUL_ESC)
    cx.font = Font(bold=True, color=BRANCO, size=10, name='Calibri')
    cx.alignment = Alignment(horizontal='center', vertical='center')
    cx.border = bdr(BRANCO)

ws_l.freeze_panes = "A3"

# Validações
cats = '"Moradia,Transporte,Alimentacao,Saude,Educacao,Lazer,Assinaturas,Comunicacao,MEI Negocios,Avivados,Parcelamentos,Outros,Salario,Renda Extra"'
subcats = '"Condominio,Energia,Internet,Financiamento Carro,Combustivel,Mercado,Restaurante,Delivery,Farmacia,Academia,Pos-Graduacao,Streaming,Apple,Photoshop,CapCut,MEI,Celular Ele,Celular Ela,Vivo,Avivados,Parcelamento Ele,Parcelamento Ela,Salario Ele,Salario Ela,Extra,Outros"'
resps = '"Ele,Ela,Casal"'
pags = '"PIX,Debito,Cartao Credito,Dinheiro,TED DOC,Boleto"'
tipos = '"Receita,Despesa"'

dv_c  = DataValidation(type='list', formula1=cats,    allow_blank=True)
dv_sc = DataValidation(type='list', formula1=subcats, allow_blank=True)
dv_r  = DataValidation(type='list', formula1=resps,   allow_blank=True)
dv_p  = DataValidation(type='list', formula1=pags,    allow_blank=True)
dv_t  = DataValidation(type='list', formula1=tipos,   allow_blank=True)

dv_c.sqref  = "C3:C3000"
dv_sc.sqref = "D3:D3000"
dv_r.sqref  = "E3:E3000"
dv_p.sqref  = "F3:F3000"
dv_t.sqref  = "H3:H3000"

for dv in [dv_c, dv_sc, dv_r, dv_p, dv_t]:
    ws_l.add_data_validation(dv)

# Dados iniciais pré-preenchidos (mês atual)
d1 = today.replace(day=1)
lancamentos = [
    (d1,"Salario - Ele",        "Salario",       "Salario Ele",        "Ele",   "TED DOC",       3600.00, "Receita"),
    (d1,"Salario - Ela",        "Salario",       "Salario Ela",        "Ela",   "TED DOC",       2500.00, "Receita"),
    (d1,"Renda Extra",          "Renda Extra",   "Extra",              "Casal", "PIX",            500.00, "Receita"),
    (d1,"Condominio",           "Moradia",       "Condominio",         "Casal", "Boleto",         426.15, "Despesa"),
    (d1,"Energia Eletrica",     "Moradia",       "Energia",            "Casal", "Debito",         220.00, "Despesa"),
    (d1,"Internet",             "Comunicacao",   "Internet",           "Casal", "Debito",         124.74, "Despesa"),
    (d1,"Streaming (Netflix etc)","Lazer",       "Streaming",          "Casal", "Debito",          40.00, "Despesa"),
    (d1,"Financiamento Carro",  "Transporte",    "Financiamento Carro","Casal", "Debito",         598.00, "Despesa"),
    (d1,"Pos-Graduacao",        "Educacao",      "Pos-Graduacao",      "Ele",   "Boleto",         115.78, "Despesa"),
    (d1,"MEI",                  "MEI Negocios",  "MEI",                "Ele",   "Boleto",          85.00, "Despesa"),
    (d1,"Avivados",             "Avivados",      "Avivados",           "Casal", "PIX",            200.00, "Despesa"),
    (d1,"Celular - Ele",        "Comunicacao",   "Celular Ele",        "Ele",   "Debito",         235.00, "Despesa"),
    (d1,"Academia",             "Saude",         "Academia",           "Ela",   "Cartao Credito",  94.90, "Despesa"),
    (d1,"Celular - Ela",        "Comunicacao",   "Celular Ela",        "Ela",   "Cartao Credito", 179.77, "Despesa"),
    (d1,"Apple (conta)",        "Assinaturas",   "Apple",              "Ele",   "Cartao Credito",   5.90, "Despesa"),
    (d1,"Adobe Photoshop",      "Assinaturas",   "Photoshop",          "Ele",   "Cartao Credito",  49.00, "Despesa"),
    (d1,"CapCut Pro",           "Assinaturas",   "CapCut",             "Ela",   "Cartao Credito",  32.00, "Despesa"),
    (d1,"Apple (conta 2)",      "Assinaturas",   "Apple",              "Ela",   "Cartao Credito",  66.00, "Despesa"),
    (d1,"Vivo",                 "Comunicacao",   "Vivo",               "Ela",   "Cartao Credito",  86.00, "Despesa"),
    (d1,"Mercado",              "Alimentacao",   "Mercado",            "Casal", "PIX",           1200.00, "Despesa"),
    (d1,"Combustivel",          "Transporte",    "Combustivel",        "Ele",   "Debito",         400.00, "Despesa"),
    (d1,"Lazer",                "Lazer",         "Outros",             "Casal", "PIX",            400.00, "Despesa"),
    (d1,"Parcelamento - Ele",   "Parcelamentos", "Parcelamento Ele",   "Ele",   "Cartao Credito", 400.00, "Despesa"),
    (d1,"Parcelamento - Ela",   "Parcelamentos", "Parcelamento Ela",   "Ela",   "Cartao Credito", 187.82, "Despesa"),
]

for i, row in enumerate(lancamentos, 3):
    H(ws_l, {i: 22})
    row_bg = CINZA_CLA if i % 2 == 0 else BRANCO
    for j, val in enumerate(row, 1):
        cx = ws_l.cell(row=i, column=j, value=val)
        cx.font = Font(size=10, name='Calibri',
                       color=(VERDE_ESC if (j==8 and val=="Receita") else
                              VERM_ESC  if (j==8 and val=="Despesa") else PRETO),
                       bold=(j==8))
        cx.alignment = Alignment(horizontal='center' if j != 2 else 'left',
                                  vertical='center')
        cx.fill = f(VERDE_CLA if (j==8 and val=="Receita") else
                    VERM_CLA  if (j==8 and val=="Despesa") else row_bg)
        cx.border = bdr()
        if j == 1:
            cx.number_format = 'DD/MM/YYYY'
        elif j == 7:
            cx.number_format = 'R$ #,##0.00'

# Formatação condicional tipo
ws_l.conditional_formatting.add('H3:H3000',
    CellIsRule(operator='equal', formula=['"Receita"'],
               font=Font(bold=True, color=VERDE_ESC, name='Calibri'),
               fill=PatternFill(fill_type='solid', fgColor=VERDE_CLA)))
ws_l.conditional_formatting.add('H3:H3000',
    CellIsRule(operator='equal', formula=['"Despesa"'],
               font=Font(bold=True, color=VERM_ESC, name='Calibri'),
               fill=PatternFill(fill_type='solid', fgColor=VERM_CLA)))

# =============================================
# ABA 3 — ORÇAMENTO MENSAL
# =============================================
ws_o = wb.create_sheet("Orcamento", 2)
ws_o.sheet_view.showGridLines = False
ws_o.sheet_properties.tabColor = VERDE_ESC

W(ws_o, [2, 26, 18, 18, 18, 16, 18, 2])
H(ws_o, {1:14, 2:42, 3:8, 4:30})

m(ws_o, 1, 2, 1, 7)
ws_o['B1'].value = "ORCAMENTO MENSAL"
ws_o['B1'].fill = f(VERDE_ESC)
ws_o['B1'].font = Font(bold=True, color=BRANCO, size=16, name='Calibri')
ws_o['B1'].alignment = Alignment(horizontal='center', vertical='center')

m(ws_o, 2, 2, 2, 4)
ws_o['B2'].value = "Mes de Referencia: " + today.strftime("%B %Y")
ws_o['B2'].fill = f(VERDE_CLA)
ws_o['B2'].font = Font(bold=True, color=VERDE_ESC, size=11, name='Calibri')
ws_o['B2'].alignment = Alignment(horizontal='center', vertical='center')

m(ws_o, 2, 5, 2, 7)
ws_o['E2'].value = "Renda Total: R$ 6.600,00"
ws_o['E2'].fill = f(AZUL_CLA)
ws_o['E2'].font = Font(bold=True, color=AZUL_ESC, size=11, name='Calibri')
ws_o['E2'].alignment = Alignment(horizontal='center', vertical='center')

# Cabeçalhos
orc_headers = ["Categoria", "Planejado (R$)", "Realizado (R$)", "Diferenca (R$)", "% Utilizado", "Status"]
for i, h in enumerate(orc_headers, 2):
    cx = ws_o.cell(row=4, column=i, value=h)
    cx.fill = f(VERDE_ESC)
    cx.font = Font(bold=True, color=BRANCO, size=10, name='Calibri')
    cx.alignment = Alignment(horizontal='center', vertical='center')
    cx.border = bdr(BRANCO)

# Categorias e valores planejados
orc_data = [
    ("Moradia",       646.15,  "Moradia"),
    ("Transporte",    998.00,  "Transporte"),
    ("Alimentacao",  1200.00,  "Alimentacao"),
    ("Saude",          94.90,  "Saude"),
    ("Educacao",      115.78,  "Educacao"),
    ("Lazer",         440.00,  "Lazer"),
    ("Assinaturas",   152.90,  "Assinaturas"),
    ("Comunicacao",   625.51,  "Comunicacao"),
    ("MEI Negocios",   85.00,  "MEI Negocios"),
    ("Avivados",      200.00,  "Avivados"),
    ("Parcelamentos", 587.82,  "Parcelamentos"),
    ("Outros",          0.00,  "Outros"),
]

for idx, (nome, planejado, cat_filter) in enumerate(orc_data):
    row = idx + 5
    H(ws_o, {row: 24})
    row_bg = CINZA_CLA if idx % 2 == 0 else BRANCO

    # Categoria
    cx = ws_o.cell(row=row, column=2, value=nome)
    cx.fill = f(row_bg)
    cx.font = Font(bold=True, size=10, color=PRETO, name='Calibri')
    cx.alignment = Alignment(horizontal='left', vertical='center')
    cx.border = bdr()

    # Planejado
    cx = ws_o.cell(row=row, column=3, value=planejado)
    cx.fill = f(row_bg)
    cx.font = Font(size=10, name='Calibri')
    cx.alignment = Alignment(horizontal='center', vertical='center')
    cx.border = bdr()
    cx.number_format = 'R$ #,##0.00'

    # Realizado (SUMPRODUCT da aba Lancamentos por categoria)
    formula_real = (
        f'=SUMPRODUCT('
        f'(MONTH(Lancamentos!$A$3:$A$3000)=MONTH(TODAY()))*'
        f'(YEAR(Lancamentos!$A$3:$A$3000)=YEAR(TODAY()))*'
        f'(Lancamentos!$H$3:$H$3000="Despesa")*'
        f'(Lancamentos!$C$3:$C$3000=B{row})*'
        f'Lancamentos!$G$3:$G$3000)'
    )
    cx = ws_o.cell(row=row, column=4, value=formula_real)
    cx.fill = f(row_bg)
    cx.font = Font(size=10, name='Calibri')
    cx.alignment = Alignment(horizontal='center', vertical='center')
    cx.border = bdr()
    cx.number_format = 'R$ #,##0.00'

    # Diferença
    cx = ws_o.cell(row=row, column=5, value=f'=C{row}-D{row}')
    cx.fill = f(row_bg)
    cx.font = Font(size=10, name='Calibri')
    cx.alignment = Alignment(horizontal='center', vertical='center')
    cx.border = bdr()
    cx.number_format = 'R$ #,##0.00'

    # Percentual
    cx = ws_o.cell(row=row, column=6, value=f'=IF(C{row}=0,0,D{row}/C{row})')
    cx.fill = f(row_bg)
    cx.font = Font(size=10, name='Calibri')
    cx.alignment = Alignment(horizontal='center', vertical='center')
    cx.border = bdr()
    cx.number_format = '0%'

    # Status
    cx = ws_o.cell(row=row, column=7,
                   value=f'=IF(C{row}=0,"---",IF(F{row}>1,"ESTOROU","OK"))')
    cx.fill = f(row_bg)
    cx.font = Font(size=10, name='Calibri')
    cx.alignment = Alignment(horizontal='center', vertical='center')
    cx.border = bdr()

# Formatação condicional: % > 100% = vermelho, > 80% = amarelo
RED_FONT  = Font(bold=True, color=BRANCO, name='Calibri')
RED_FILL  = PatternFill(fill_type='solid', fgColor=VERM_ESC)
YEL_FONT  = Font(bold=True, color=PRETO,  name='Calibri')
YEL_FILL  = PatternFill(fill_type='solid', fgColor='F9E79F')
GRN_FONT  = Font(bold=True, color=VERDE_ESC, name='Calibri')
GRN_FILL  = PatternFill(fill_type='solid', fgColor=VERDE_CLA)

ws_o.conditional_formatting.add(f'F5:F{5+len(orc_data)-1}',
    CellIsRule(operator='greaterThan', formula=['1'], font=RED_FONT, fill=RED_FILL))
ws_o.conditional_formatting.add(f'F5:F{5+len(orc_data)-1}',
    CellIsRule(operator='between', formula=['0.8', '1'], font=YEL_FONT, fill=YEL_FILL))
ws_o.conditional_formatting.add(f'F5:F{5+len(orc_data)-1}',
    CellIsRule(operator='lessThan', formula=['0.8'], font=GRN_FONT, fill=GRN_FILL))

ws_o.conditional_formatting.add(f'G5:G{5+len(orc_data)-1}',
    CellIsRule(operator='equal', formula=['"ESTOROU"'], font=RED_FONT, fill=RED_FILL))
ws_o.conditional_formatting.add(f'G5:G{5+len(orc_data)-1}',
    CellIsRule(operator='equal', formula=['"OK"'], font=GRN_FONT, fill=GRN_FILL))

# Total
total_row = 5 + len(orc_data)
H(ws_o, {total_row: 28})
for col, val in [(2, "TOTAL ORCADO"), (3, f'=SUM(C5:C{total_row-1})'),
                  (4, f'=SUM(D5:D{total_row-1})'), (5, f'=C{total_row}-D{total_row}'),
                  (6, f'=IF(C{total_row}=0,0,D{total_row}/C{total_row})'), (7, "")]:
    cx = ws_o.cell(row=total_row, column=col, value=val)
    cx.fill = f(AZUL_ESC)
    cx.font = Font(bold=True, color=BRANCO, size=10, name='Calibri')
    cx.alignment = Alignment(horizontal='center', vertical='center')
    cx.border = bdr(BRANCO)
    if col in [3, 4, 5]:
        cx.number_format = 'R$ #,##0.00'
    elif col == 6:
        cx.number_format = '0%'

# Renda e sobra
renda_row = total_row + 2
H(ws_o, {renda_row: 26, renda_row+1: 26})

for col, val, bg in [
    (2, "Renda Total do Mes", AZUL_ESC),
    (3, '=SUMPRODUCT((MONTH(Lancamentos!$A$3:$A$3000)=MONTH(TODAY()))*(YEAR(Lancamentos!$A$3:$A$3000)=YEAR(TODAY()))*(Lancamentos!$H$3:$H$3000="Receita")*Lancamentos!$G$3:$G$3000)', AZUL_ESC),
    (4, "", AZUL_ESC), (5,"",AZUL_ESC), (6,"",AZUL_ESC), (7,"",AZUL_ESC)]:
    cx = ws_o.cell(row=renda_row, column=col, value=val)
    cx.fill = f(bg)
    cx.font = Font(bold=True, color=BRANCO, size=10, name='Calibri')
    cx.alignment = Alignment(horizontal='center' if col != 2 else 'left', vertical='center')
    if col == 3:
        cx.number_format = 'R$ #,##0.00'
    cx.border = bdr(BRANCO)

for col, val, bg in [
    (2, "Sobra / Disponivel", VERDE_ESC),
    (3, f'=C{renda_row}-C{total_row}', VERDE_ESC),
    (4, "", VERDE_ESC),(5,"",VERDE_ESC),(6,"",VERDE_ESC),(7,"",VERDE_ESC)]:
    cx = ws_o.cell(row=renda_row+1, column=col, value=val)
    cx.fill = f(bg)
    cx.font = Font(bold=True, color=BRANCO, size=11, name='Calibri')
    cx.alignment = Alignment(horizontal='center' if col != 2 else 'left', vertical='center')
    if col == 3:
        cx.number_format = 'R$ #,##0.00'
    cx.border = bdr(BRANCO)

# =============================================
# ABA 4 — RESERVA DE EMERGÊNCIA
# =============================================
ws_r = wb.create_sheet("Reserva", 3)
ws_r.sheet_view.showGridLines = False
ws_r.sheet_properties.tabColor = ROXO_ESC

W(ws_r, [2, 34, 22, 2])
H(ws_r, {1:14, 2:42, 3:12})

m(ws_r, 1, 2, 1, 3)
ws_r['B1'].value = "RESERVA DE EMERGENCIA"
ws_r['B1'].fill = f(ROXO_ESC)
ws_r['B1'].font = Font(bold=True, color=BRANCO, size=16, name='Calibri')
ws_r['B1'].alignment = Alignment(horizontal='center', vertical='center')

m(ws_r, 2, 2, 2, 3)
ws_r['B2'].value = "Sua seguranca financeira comeca aqui. Atualize o campo amarelo sempre que depositar."
ws_r['B2'].fill = f(ROXO_CLA)
ws_r['B2'].font = Font(color=ROXO_ESC, size=10, italic=True, name='Calibri')
ws_r['B2'].alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

campos = [
    (4,  "Custo de Vida Mensal (R$)",          5146.06, 'R$ #,##0.00', False, ROXO_CLA,  PRETO),
    (5,  "Meses de Protecao Desejados",         6,       '0',           False, ROXO_CLA,  PRETO),
    (6,  "META TOTAL DA RESERVA (R$)",          "=B4*B5",'R$ #,##0.00', False, ROXO_ESC,  BRANCO),
    (7,  "VALOR JA ACUMULADO (R$)  <- EDITE AQUI!", 0,  'R$ #,##0.00', True,  AMAR_CLA,  LARAN),
    (8,  "Falta para Completar (R$)",           "=B6-B7",'R$ #,##0.00', False, VERM_CLA,  VERM_ESC),
    (9,  "Percentual Concluido (%)",  "=IF(B6=0,0,B7/B6)",'0.0%',      False, VERDE_CLA, VERDE_ESC),
    (10, "Meses de Vida Protegidos (atual)","=IF(B4=0,0,B7/B4)",'0.0', False, AZUL_CLA,  AZUL_ESC),
    (11, "Aporte Mensal Sugerido (R$)",          1454.00,'R$ #,##0.00', True,  AMAR_CLA,  LARAN),
    (12, "Previsao de Conclusao",
         '=IFERROR(TEXT(TODAY()+((B6-B7)/B11)*30,"MMM/YYYY"),"---")',
         '@', False, AZUL_CLA, AZUL_ESC),
]

for (row, label, val, fmt, editable, cell_bg, cell_fg) in campos:
    H(ws_r, {row: 30})

    cx = ws_r.cell(row=row, column=2, value=label)
    cx.fill = f(CINZA_CLA if not editable else AMAR_CLA)
    cx.font = Font(bold=True, size=10, color=LARAN if editable else PRETO, name='Calibri')
    cx.alignment = Alignment(horizontal='left', vertical='center', indent=1)
    cx.border = bdr()

    cx = ws_r.cell(row=row, column=3, value=val)
    cx.fill = f(cell_bg)
    cx.font = Font(bold=True, size=12, color=cell_fg, name='Calibri')
    cx.alignment = Alignment(horizontal='center', vertical='center')
    cx.border = bdr_thick(ROXO_ESC) if editable else bdr()
    cx.number_format = fmt

# Barra de progresso visual (simulada)
H(ws_r, {14: 12, 15: 36, 16: 12})
m(ws_r, 14, 2, 14, 3)
ws_r['B14'].value = "PROGRESSO"
ws_r['B14'].fill = f(ROXO_ESC)
ws_r['B14'].font = Font(bold=True, color=BRANCO, size=9, name='Calibri')
ws_r['B14'].alignment = Alignment(horizontal='center', vertical='center')

m(ws_r, 15, 2, 15, 3)
cx = ws_r['B15']
cx.value = ('=IFERROR('
            '"["&REPT("|",ROUND(B9*30,0))&REPT(".",30-ROUND(B9*30,0))&"]  "&TEXT(B9,"0%")&" concluido",'
            '"Nenhum valor acumulado ainda")')
cx.fill = f(VERDE_CLA)
cx.font = Font(bold=True, color=VERDE_ESC, size=11, name='Calibri')
cx.alignment = Alignment(horizontal='center', vertical='center')
cx.border = bdr_thick(VERDE_ESC)

# Dica
H(ws_r, {18: 22})
m(ws_r, 18, 2, 18, 3)
ws_r['B18'].value = "  Onde guardar: Tesouro Selic ou CDB 100% CDI com liquidez diaria. Nunca misture com conta corrente."
ws_r['B18'].fill = f(AMAR_CLA)
ws_r['B18'].font = Font(color=LARAN, size=9, italic=True, name='Calibri')
ws_r['B18'].alignment = Alignment(horizontal='left', vertical='center')

# =============================================
# ABA 5 — DÍVIDAS
# =============================================
ws_div = wb.create_sheet("Dividas", 4)
ws_div.sheet_view.showGridLines = False
ws_div.sheet_properties.tabColor = VERM_ESC

W(ws_div, [2, 22, 14, 16, 16, 14, 12, 16, 14, 2])
H(ws_div, {1:14, 2:42, 3:12, 4:30})

m(ws_div, 1, 2, 1, 9)
ws_div['B1'].value = "DIVIDAS E PARCELAMENTOS"
ws_div['B1'].fill = f(VERM_ESC)
ws_div['B1'].font = Font(bold=True, color=BRANCO, size=16, name='Calibri')
ws_div['B1'].alignment = Alignment(horizontal='center', vertical='center')

div_headers = ["Credor", "Tipo", "Saldo Total (R$)", "Parcela Mensal (R$)",
               "Parcelas Rest.", "Juros (% a.m.)", "Termino", "Prioridade"]
for i, h in enumerate(div_headers, 2):
    cx = ws_div.cell(row=4, column=i, value=h)
    cx.fill = f(VERM_ESC)
    cx.font = Font(bold=True, color=BRANCO, size=9, name='Calibri')
    cx.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cx.border = bdr(BRANCO)

# Única dívida conhecida
H(ws_div, {5: 26})
div_row = [
    "Financiamento Carro", "Financiamento", None, 598.00, None, None, None, "Alta"
]
for j, val in enumerate(div_row, 2):
    cx = ws_div.cell(row=5, column=j, value=val)
    cx.fill = f(VERM_CLA)
    cx.font = Font(size=10, name='Calibri')
    cx.alignment = Alignment(horizontal='center', vertical='center')
    cx.border = bdr()
    if j == 4:
        cx.number_format = 'R$ #,##0.00'
    if j == 3:
        cx.number_format = 'R$ #,##0.00'

# Linhas em branco para novas dívidas
for row in range(6, 16):
    H(ws_div, {row: 22})
    for col in range(2, 10):
        cx = ws_div.cell(row=row, column=col)
        cx.fill = f(CINZA_CLA if row % 2 == 0 else BRANCO)
        cx.border = bdr()
        if col in [3, 4]:
            cx.number_format = 'R$ #,##0.00'

# Total dívidas
H(ws_div, {16: 28})
m(ws_div, 16, 2, 16, 3)
ws_div.cell(row=16, column=2, value="TOTAL RESTANTE (R$)").fill = f(AZUL_ESC)
ws_div['B16'].font = Font(bold=True, color=BRANCO, size=10, name='Calibri')
ws_div['B16'].alignment = Alignment(horizontal='center', vertical='center')

cx = ws_div.cell(row=16, column=4, value='=IFERROR(SUM(D5:D15),0)')
cx.fill = f(AZUL_ESC)
cx.font = Font(bold=True, color=BRANCO, size=10, name='Calibri')
cx.number_format = 'R$ #,##0.00'
cx.alignment = Alignment(horizontal='center', vertical='center')

# Dica
H(ws_div, {18: 22})
m(ws_div, 18, 2, 18, 9)
ws_div['B18'].value = "  Preencha: Saldo Total, Parcelas Restantes, Juros e Data de Termino para cada divida."
ws_div['B18'].fill = f(AMAR_CLA)
ws_div['B18'].font = Font(color=LARAN, size=9, italic=True, name='Calibri')
ws_div['B18'].alignment = Alignment(horizontal='left', vertical='center')

# =============================================
# ABA 6 — METAS
# =============================================
ws_m = wb.create_sheet("Metas", 5)
ws_m.sheet_view.showGridLines = False
ws_m.sheet_properties.tabColor = LARAN

W(ws_m, [2, 26, 18, 14, 18, 14, 18, 2])
H(ws_m, {1:14, 2:42, 3:12, 4:30})

m(ws_m, 1, 2, 1, 7)
ws_m['B1'].value = "METAS FINANCEIRAS"
ws_m['B1'].fill = f(LARAN)
ws_m['B1'].font = Font(bold=True, color=BRANCO, size=16, name='Calibri')
ws_m['B1'].alignment = Alignment(horizontal='center', vertical='center')

met_headers = ["Meta", "Valor Alvo (R$)", "Prazo", "Valor Atual (R$)", "Falta (R$)", "% Concluido"]
for i, h in enumerate(met_headers, 2):
    cx = ws_m.cell(row=4, column=i, value=h)
    cx.fill = f(LARAN)
    cx.font = Font(bold=True, color=BRANCO, size=10, name='Calibri')
    cx.alignment = Alignment(horizontal='center', vertical='center')
    cx.border = bdr(BRANCO)

metas = [
    ("Reserva de Emergencia (6 meses)", "=Reserva!B6", "Dez/2027", "=Reserva!B7"),
    ("Entrada - Imovel Proprio",         None,          "Jun/2029",  0),
    ("",                                 None,          "",          0),
    ("",                                 None,          "",          0),
]

for idx, (nome, alvo, prazo, atual) in enumerate(metas):
    row = idx + 5
    H(ws_m, {row: 30})
    row_bg = LARAN_CLA if idx % 2 == 0 else BRANCO

    vals = [nome, alvo, prazo, atual,
            f'=IF(OR(C{row}=0,C{row}=""),0,C{row}-E{row})' if alvo else "",
            f'=IF(OR(C{row}=0,C{row}=""),0,E{row}/C{row})' if alvo else ""]

    for j, val in enumerate(vals, 2):
        cx = ws_m.cell(row=row, column=j, value=val)
        cx.fill = f(row_bg)
        cx.font = Font(size=10, name='Calibri',
                       bold=(j==2), color=PRETO)
        cx.alignment = Alignment(horizontal='left' if j==2 else 'center', vertical='center')
        cx.border = bdr()
        if j in [3, 4, 5, 6] and isinstance(alvo, (int, float, str)) and alvo:
            if j in [3, 4, 5]:
                cx.number_format = 'R$ #,##0.00'
            elif j == 7:
                cx.number_format = '0%'
        if j == 7:
            cx.number_format = '0%'

# Formatação condicional % metas
ws_m.conditional_formatting.add(f'G5:G{5+len(metas)-1}',
    CellIsRule(operator='greaterThanOrEqual', formula=['1'], font=GRN_FONT, fill=GRN_FILL))
ws_m.conditional_formatting.add(f'G5:G{5+len(metas)-1}',
    CellIsRule(operator='between', formula=['0.5', '0.99'], font=YEL_FONT, fill=YEL_FILL))
ws_m.conditional_formatting.add(f'G5:G{5+len(metas)-1}',
    CellIsRule(operator='lessThan', formula=['0.5'], font=RED_FONT, fill=RED_FILL))

# Dica metas
H(ws_m, {10: 22})
m(ws_m, 10, 2, 10, 7)
ws_m['B10'].value = "  Atualize o campo 'Valor Atual' sempre que fizer um deposito ou guardar dinheiro para a meta."
ws_m['B10'].fill = f(AMAR_CLA)
ws_m['B10'].font = Font(color=LARAN, size=9, italic=True, name='Calibri')
ws_m['B10'].alignment = Alignment(horizontal='left', vertical='center')

# =============================================
# ABA AUXILIAR (listas ocultas — não exibida)
# =============================================
ws_aux = wb.create_sheet("Aux", 6)
ws_aux.sheet_state = 'hidden'
listas = {
    'A': ["Moradia","Transporte","Alimentacao","Saude","Educacao","Lazer",
          "Assinaturas","Comunicacao","MEI Negocios","Avivados","Parcelamentos","Outros","Salario","Renda Extra"],
    'B': ["Condominio","Energia","Internet","Financiamento Carro","Combustivel",
          "Mercado","Restaurante","Delivery","Farmacia","Academia","Pos-Graduacao",
          "Streaming","Apple","Photoshop","CapCut","MEI","Celular Ele","Celular Ela",
          "Vivo","Avivados","Parcelamento Ele","Parcelamento Ela","Salario Ele","Salario Ela","Extra","Outros"],
}
for col, items in listas.items():
    for i, item in enumerate(items, 1):
        ws_aux.cell(row=i, column=ord(col)-64, value=item)

# =============================================
# AJUSTES FINAIS
# =============================================
wb.active = ws_d  # Dashboard como primeira aba visível

output_path = "/home/user/Teste/Controle_Financeiro_Familiar.xlsx"
wb.save(output_path)
print(f"Planilha gerada: {output_path}")
